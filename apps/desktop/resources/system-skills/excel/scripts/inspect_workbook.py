#!/usr/bin/env python3
"""输出 Excel 工作簿的公式、缓存、宏、外链和兼容风险摘要。"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}
RISKY_FUNCTIONS = ("XLOOKUP", "XMATCH", "FILTER", "SORT", "UNIQUE", "SEQUENCE")
PROFILE_MAX_ROWS = 2000
PROFILE_MAX_COLUMNS = 64
HEADER_SCAN_ROWS = 12


def value_kind(value) -> str:
    if value is None or value == "":
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def build_data_profile(sheet) -> dict:
    columns: dict[int, dict] = {}
    header_candidates: list[dict] = []
    max_row = min(sheet.max_row, PROFILE_MAX_ROWS)
    max_column = min(sheet.max_column, PROFILE_MAX_COLUMNS)
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_column),
        start=1,
    ):
        header_values: list[str] = []
        header_text_count = 0
        for column_index, cell in enumerate(row, start=1):
            value = cell.value
            kind = value_kind(value)
            profile = columns.setdefault(
                column_index,
                {"type_counts": Counter(), "zero_count": 0, "samples": []},
            )
            profile["type_counts"][kind] += 1
            if kind == "number" and value == 0:
                profile["zero_count"] += 1
            if kind not in {"blank", "formula"} and len(profile["samples"]) < 3:
                sample = str(value)[:120]
                if sample not in profile["samples"]:
                    profile["samples"].append(sample)
            if row_index <= HEADER_SCAN_ROWS and value not in (None, ""):
                rendered = str(value)[:120]
                header_values.append(rendered)
                if kind == "text":
                    header_text_count += 1
        if row_index <= HEADER_SCAN_ROWS and header_values:
            header_candidates.append(
                {
                    "row": row_index,
                    "score": header_text_count * 2 + len(set(header_values)),
                    "values": header_values[:20],
                }
            )
    rendered_columns: list[dict] = []
    mixed_type_columns: list[str] = []
    for column_index, profile in sorted(columns.items()):
        counts = dict(profile["type_counts"])
        observed = [
            kind
            for kind in ("number", "date", "boolean", "text")
            if counts.get(kind, 0)
        ]
        letter = get_column_letter(column_index)
        if len(observed) > 1:
            mixed_type_columns.append(letter)
        rendered_columns.append(
            {
                "column": letter,
                "type_counts": counts,
                "zero_count": profile["zero_count"],
                "samples": profile["samples"],
            }
        )
    header_candidates.sort(key=lambda item: (-item["score"], item["row"]))
    return {
        "profiled_rows": max_row,
        "profiled_columns": max_column,
        "header_candidates": header_candidates[:3],
        "mixed_type_columns": mixed_type_columns,
        "columns": rendered_columns,
        "truncated": sheet.max_row > max_row or sheet.max_column > max_column,
    }


def inspect(path: Path) -> dict:
    keep_vba = path.suffix.lower() == ".xlsm"
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
    formulas_book = load_workbook(path, data_only=False, read_only=False, keep_vba=keep_vba)
    values_book = load_workbook(path, data_only=True, read_only=False, keep_vba=keep_vba)
    formula_count = 0
    missing_cache = 0
    errors: list[str] = []
    risky: list[dict[str, str]] = []
    external_formula_cells: list[str] = []
    sheets: list[dict] = []
    for sheet_name in formulas_book.sheetnames:
        formula_sheet = formulas_book[sheet_name]
        value_sheet = values_book[sheet_name]
        sheets.append(
            {
                "name": sheet_name,
                "state": formula_sheet.sheet_state,
                "max_row": formula_sheet.max_row,
                "max_column": formula_sheet.max_column,
                "merged_ranges": len(formula_sheet.merged_cells.ranges),
                "data_profile": build_data_profile(formula_sheet),
            }
        )
        for row in formula_sheet.iter_rows():
            for cell in row:
                value = cell.value
                cached = value_sheet[cell.coordinate].value
                location = f"{sheet_name}!{cell.coordinate}"
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                    if cached is None:
                        missing_cache += 1
                    upper = value.upper()
                    for function in RISKY_FUNCTIONS:
                        if re.search(rf"(?<![A-Z0-9_.])(?:_XLFN\.)?{function}\s*\(", upper):
                            risky.append({"cell": location, "function": function})
                    if re.search(r"'?\[[^\]]+\]", value):
                        external_formula_cells.append(location)
                if isinstance(cached, str) and cached in ERROR_VALUES:
                    errors.append(f"{location}:{cached}")
    formulas_book.close()
    values_book.close()
    error_counts = Counter(item.rsplit(":", 1)[1] for item in errors)
    external_parts = sorted(name for name in names if name.startswith("xl/externalLinks/") and name.endswith(".xml"))
    return {
        "status": "success",
        "file": str(path),
        "sheets": sheets,
        "formula_count": formula_count,
        "missing_formula_cache": missing_cache,
        "formula_errors": errors[:200],
        "formula_error_counts": dict(error_counts),
        "formula_errors_truncated": max(0, len(errors) - 200),
        "has_vba": "xl/vbaProject.bin" in names,
        "external_link_parts": external_parts,
        "external_formula_cells": external_formula_cells[:200],
        "risky_formulas": risky[:200],
        "risks_truncated": max(0, len(risky) - 200),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="检查 Excel 工作簿")
    parser.add_argument("path", type=Path)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    try:
        result = inspect(args.path.resolve())
        print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else json.dumps(result, ensure_ascii=False))
        return 0
    except Exception as error:
        result = {"status": "failed", "error": str(error)}
        print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else result["error"])
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
