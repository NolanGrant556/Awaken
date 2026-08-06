#!/usr/bin/env python3
"""使用独立 LibreOffice profile 重算 XLSX，并阻止未写回却报告成功。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

sys.dont_write_bytecode = True

from openpyxl import load_workbook

from office_runtime import find_soffice, office_env, temporary_directory


ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NULL!", "#NUM!"}
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
Sub AwakenRecalculateAndSave()
  ThisComponent.calculateAll()
  ThisComponent.store()
  ThisComponent.close(True)
End Sub
</script:module>
"""


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_risks(path: Path) -> tuple[int, list[str]]:
    with zipfile.ZipFile(path) as package:
        external = sorted(name for name in package.namelist() if name.startswith("xl/externalLinks/") and name.endswith(".xml"))
    workbook = load_workbook(path, data_only=False, read_only=False)
    formulas = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
    workbook.close()
    return formulas, external


def scan_errors(path: Path) -> dict:
    values = load_workbook(path, data_only=True, read_only=False)
    formulas = load_workbook(path, data_only=False, read_only=False)
    locations: list[str] = []
    formula_count = 0
    for sheet in formulas.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                    locations.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    values.close()
    formulas.close()
    counts = Counter(item.rsplit(":", 1)[1] for item in locations)
    return {
        "status": "success" if not locations else "errors_found",
        "total_formulas": formula_count,
        "total_errors": len(locations),
        "error_summary": dict(counts),
        "error_locations": locations[:200],
        "error_locations_truncated": max(0, len(locations) - 200),
    }


def recalculate(input_path: Path, output_path: Path, timeout: int) -> dict:
    if input_path.suffix.lower() != ".xlsx":
        raise ValueError("recalc.py 只接受 .xlsx；宏和模板文件需要专门路线")
    if not input_path.is_file():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")
    formulas, external = workbook_risks(input_path)
    if external:
        raise ValueError(
            "工作簿包含外部工作簿链接，LibreOffice 重算可能破坏链接和缓存；已拒绝。"
            f" parts={external[:20]}"
        )
    executable = find_soffice()
    if executable is None:
        raise FileNotFoundError("未找到 LibreOffice")
    with temporary_directory("awaken-excel-recalc-") as root:
        profile = root / "profile"
        working = root / "working.xlsx"
        shutil.copy2(input_path, working)
        init = subprocess.run(
            [str(executable), "--headless", "--terminate_after_init", f"-env:UserInstallation={profile.resolve().as_uri()}"],
            capture_output=True,
            text=True,
            timeout=min(timeout, 30),
            check=False,
            env=office_env(),
        )
        if init.returncode != 0:
            raise RuntimeError((init.stderr or init.stdout or "LibreOffice profile 初始化失败").strip())
        macro_dir = profile / "user" / "basic" / "Standard"
        macro_dir.mkdir(parents=True, exist_ok=True)
        (macro_dir / "Module1.xba").write_text(MACRO, encoding="utf-8")
        before_hash = digest(working)
        command = [
            str(executable),
            "--headless",
            "--nologo",
            "--norestore",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "vnd.sun.star.script:Standard.Module1.AwakenRecalculateAndSave?language=Basic&location=application",
            str(working.resolve()),
        ]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False, env=office_env())
        except subprocess.TimeoutExpired as error:
            raise RuntimeError(f"LibreOffice 在 {timeout}s 后超时，工作簿没有被确认重算") from error
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or f"LibreOffice 退出码 {result.returncode}").strip())
        after_hash = digest(working)
        if formulas and before_hash == after_hash:
            raise RuntimeError("工作簿包含公式，但 LibreOffice 没有真正写回文件")
        scan = scan_errors(working)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = output_path.with_name(f".{output_path.name}.awaken-tmp")
        if temporary.exists():
            raise RuntimeError(f"临时输出已存在: {temporary}")
        shutil.copy2(working, temporary)
        os.replace(temporary, output_path)
    return {**scan, "output": str(output_path), "rewritten": before_hash != after_hash}


def main() -> int:
    parser = argparse.ArgumentParser(description="重算 XLSX 公式并扫描错误")
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--timeout", type=int, default=90)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    input_path = args.input.resolve()
    output_path = args.output.resolve() if args.output else input_path
    try:
        result = recalculate(input_path, output_path, args.timeout)
        print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else json.dumps(result, ensure_ascii=False))
        return 0 if result["status"] == "success" else 2
    except Exception as error:
        result = {"status": "failed", "error": str(error)}
        print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else result["error"])
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
