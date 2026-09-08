from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path


REPO = Path(__file__).resolve().parents[2]
SYSTEM_SKILLS = REPO / "apps" / "desktop" / "resources" / "system-skills"


def load_script(skill: str, script: str, module_name: str):
    script_dir = SYSTEM_SKILLS / skill / "scripts"
    sys.path.insert(0, str(script_dir))
    try:
        spec = importlib.util.spec_from_file_location(module_name, script_dir / script)
        if spec is None or spec.loader is None:
            raise RuntimeError(f"cannot load {script}")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.pop(0)


def write_package(path: Path, files: dict[str, str | bytes]) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as package:
        for name, content in files.items():
            package.writestr(name, content)


def minimal_docx_files() -> dict[str, str]:
    return {
        "[Content_Types].xml": """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>""",
        "_rels/.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>""",
        "word/document.xml": """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>第一</w:t></w:r><w:r><w:t>段</w:t></w:r></w:p></w:body></w:document>""",
    }


def analyzable_docx_files() -> dict[str, str | bytes]:
    files: dict[str, str | bytes] = minimal_docx_files()
    files["word/styles.xml"] = """<?xml version="1.0" encoding="UTF-8"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:pPr><w:outlineLvl w:val="0"/></w:pPr></w:style>
</w:styles>"""
    files["word/_rels/document.xml.rels"] = """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdSafe" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/a b" TargetMode="External"/>
  <Relationship Id="rIdUnsafe" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="javascript:alert(1)" TargetMode="External"/>
  <Relationship Id="rIdPackage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="media/image1.png"/>
</Relationships>"""
    files["word/document.xml"] = """<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">
 <w:body>
  <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>分析标题</w:t></w:r></w:p>
  <w:p><w:r><w:t>访问</w:t></w:r><w:hyperlink r:id="rIdSafe"><w:r><w:t>官网</w:t></w:r></w:hyperlink></w:p>
  <w:p><w:hyperlink r:id="rIdUnsafe"><w:r><w:t>危险链接文本</w:t></w:r></w:hyperlink></w:p>
  <w:p><w:hyperlink r:id="rIdPackage"><w:r><w:t>包内关系文本</w:t></w:r></w:hyperlink></w:p>
  <w:tbl><w:tr><w:tc><w:p><w:r><w:t>姓名</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>金额</w:t></w:r></w:p></w:tc></w:tr><w:tr><w:tc><w:p><w:r><w:t>甲</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>100</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
  <w:p><w:r><w:t>带脚注</w:t><w:footnoteReference w:id="2"/></w:r></w:p>
  <w:p><w:r><w:t>带批注</w:t><w:commentReference w:id="3"/></w:r></w:p>
  <w:p><w:r><w:drawing><wp:inline><wp:docPr id="1" name="Chart" descr="季度趋势图"/></wp:inline></w:drawing></w:r></w:p>
 </w:body>
</w:document>"""
    files["word/footnotes.xml"] = """<?xml version="1.0" encoding="UTF-8"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="2"><w:p><w:r><w:t>脚注正文</w:t></w:r></w:p></w:footnote></w:footnotes>"""
    files["word/comments.xml"] = """<?xml version="1.0" encoding="UTF-8"?>
<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:comment w:id="3" w:author="审阅人"><w:p><w:r><w:t>批注意见</w:t></w:r></w:p></w:comment></w:comments>"""
    files["word/media/image1.png"] = b"not-decoded-by-reader"
    return files


def minimal_pptx_files() -> dict[str, str]:
    return {
        "[Content_Types].xml": """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
</Types>""",
        "_rels/.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>""",
        "ppt/presentation.xml": """<?xml version="1.0" encoding="UTF-8"?>
<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>""",
        "ppt/_rels/presentation.xml.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
</Relationships>""",
        "ppt/slides/slide1.xml": """<?xml version="1.0" encoding="UTF-8"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cSld><p:spTree><p:sp><p:nvSpPr><p:cNvPr id="1" name="Title"/></p:nvSpPr></p:sp></p:spTree></p:cSld></p:sld>""",
    }


class SkillResourceRoutingTests(unittest.TestCase):
    def test_bundled_markdown_references_are_reachable_and_stay_in_package(self):
        import re
        from urllib.parse import unquote
        for root in SYSTEM_SKILLS.iterdir():
            if not (root / "SKILL.md").is_file():
                continue
            visited = set()
            pending = [root / "SKILL.md"]
            while pending:
                path = pending.pop().resolve()
                if path in visited:
                    continue
                visited.add(path)
                for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", path.read_text(encoding="utf-8")):
                    target = unquote(target.strip("<>").split("#", 1)[0])
                    if not target or re.match(r"[a-zA-Z]+:", target):
                        continue
                    destination = (path.parent / target).resolve()
                    self.assertTrue(destination.is_relative_to(root.resolve()), (path, target))
                    self.assertTrue(destination.is_file(), (path, target))
                    if destination.suffix == ".md":
                        pending.append(destination)
            for reference in (root / "references").glob("*.md"):
                self.assertIn(reference.resolve(), visited, f"unrouted reference: {reference}")


class WordScriptTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.validator = load_script("word", "validate_docx.py", "awaken_word_validate_test")
        cls.package = load_script("word", "ooxml_package.py", "awaken_word_package_test")
        cls.merger = load_script("word", "merge_runs.py", "awaken_word_merge_test")
        cls.acceptor = load_script("word", "accept_changes.py", "awaken_word_accept_test")
        cls.reader = load_script("word", "read_docx.py", "awaken_word_read_test")
        cls.doctor = load_script("word", "doctor.py", "awaken_word_doctor_test")

    def test_valid_docx_and_relationship_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            valid = root / "valid.docx"
            write_package(valid, minimal_docx_files())
            self.assertEqual([], self.validator.inspect(valid))

            broken_files = minimal_docx_files()
            broken_files["_rels/.rels"] = broken_files["_rels/.rels"].replace(
                "word/document.xml", "word/missing.xml"
            )
            broken = root / "broken.docx"
            write_package(broken, broken_files)
            self.assertTrue(
                any("关系目标不存在" in issue for issue in self.validator.inspect(broken))
            )

    def test_safe_unpack_rejects_path_escape(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            archive = root / "unsafe.docx"
            files = minimal_docx_files()
            files["../escape.txt"] = "escape"
            write_package(archive, files)
            with self.assertRaises(ValueError):
                self.package.safe_extract(archive, root / "unpacked")

    def test_merge_only_plain_identical_runs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            unpacked = root / "unpacked"
            unpacked.mkdir()
            package = root / "input.docx"
            write_package(package, minimal_docx_files())
            self.package.safe_extract(package, unpacked)
            merged = self.merger.merge_document(unpacked / "word" / "document.xml")
            self.assertEqual(1, merged)
            text = (unpacked / "word" / "document.xml").read_text(encoding="utf-8")
            self.assertIn("第一段", text)

    def test_tracked_changes_cover_headers_and_other_word_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "header-revision.docx"
            files = minimal_docx_files()
            files["word/header1.xml"] = """<?xml version="1.0" encoding="UTF-8"?>
<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:ins w:id="1"><w:r><w:t>页眉修订</w:t></w:r></w:ins></w:p></w:hdr>"""
            write_package(path, files)
            self.assertEqual(1, self.acceptor.tracked_change_count(path))
            self.assertTrue(
                any("word/header1.xml" in issue and "w:author" in issue for issue in self.validator.inspect(path))
            )

    def test_stdlib_reader_extracts_document_structure_and_warnings(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "analysis.docx"
            write_package(path, analyzable_docx_files())
            result = self.reader.read_docx(path)

            self.assertEqual("success_with_warnings", result["status"])
            self.assertIn("# 分析标题", result["markdown"])
            self.assertIn("[官网](https://example.com/a%20b)", result["markdown"])
            self.assertIn("危险链接文本", result["markdown"])
            self.assertNotIn("javascript:", result["markdown"])
            self.assertIn("包内关系文本", result["markdown"])
            self.assertNotIn("media/image1.png", result["markdown"])
            self.assertIn("| 姓名 | 金额 |", result["markdown"])
            self.assertIn("[^脚注-2]: 脚注正文", result["markdown"])
            self.assertIn("批注-3（审阅人）: 批注意见", result["markdown"])
            self.assertIn("[图片: 季度趋势图", result["markdown"])
            table = next(block for block in result["blocks"] if block["type"] == "table")
            self.assertEqual([["姓名", "金额"], ["甲", "100"]], table["rows"])
            warning_codes = {warning["code"] for warning in result["warnings"]}
            self.assertIn("unsafe_hyperlink_not_linked", warning_codes)
            self.assertIn("package_relationship_not_linked", warning_codes)
            self.assertIn("media_not_interpreted", warning_codes)
            self.assertEqual(1, result["stats"]["tables"])
            self.assertEqual(1, result["stats"]["footnotes"])
            self.assertEqual(1, result["stats"]["comments"])

    def test_stdlib_reader_cli_emits_json_without_external_dependencies(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "analysis.docx"
            write_package(path, analyzable_docx_files())
            script = SYSTEM_SKILLS / "word" / "scripts" / "read_docx.py"
            completed = subprocess.run(
                [sys.executable, "-B", str(script), str(path), "--format", "json"],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(0, completed.returncode, completed.stderr)
            result = json.loads(completed.stdout)
            self.assertEqual("success_with_warnings", result["status"])
            self.assertIn("分析标题", result["markdown"])

    def test_stdlib_reader_rejects_unsafe_package_member(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "unsafe.docx"
            files = minimal_docx_files()
            files["../escape.xml"] = "<escape/>"
            write_package(path, files)
            with self.assertRaises(self.reader.DocxReadError):
                self.reader.read_docx(path)

    def test_stdlib_reader_checks_xml_size_before_read_and_wraps_runtime_error(self):
        class Info:
            def __init__(self, size: int):
                self.file_size = size

        class FakePackage:
            def __init__(self, size: int, error: Exception | None = None):
                self.info = Info(size)
                self.error = error
                self.read_called = False

            def getinfo(self, _part: str):
                return self.info

            def read(self, _info):
                self.read_called = True
                if self.error:
                    raise self.error
                return b"<root/>"

        oversized = FakePackage(self.reader.MAX_XML_BYTES + 1)
        with self.assertRaises(self.reader.DocxReadError):
            self.reader.parse_part(oversized, "word/document.xml")
        self.assertFalse(oversized.read_called)

        encrypted = FakePackage(128, RuntimeError("encrypted file"))
        with self.assertRaises(self.reader.DocxReadError) as raised:
            self.reader.parse_part(encrypted, "word/document.xml")
        self.assertIn("无法读取", str(raised.exception))

    def test_doctor_reports_routes_independently(self):
        original_command_path = self.doctor.command_path
        original_libreoffice_path = self.doctor.libreoffice_path
        original_module_available = self.doctor.module_available
        original_node_module_available = self.doctor.node_module_available
        try:
            self.doctor.command_path = lambda *_names: None
            self.doctor.libreoffice_path = lambda: None
            self.doctor.module_available = lambda _name: False
            self.doctor.node_module_available = lambda _name, _root=None: False
            result = self.doctor.inspect()
            read_only = self.doctor.inspect("read")
        finally:
            self.doctor.command_path = original_command_path
            self.doctor.libreoffice_path = original_libreoffice_path
            self.doctor.module_available = original_module_available
            self.doctor.node_module_available = original_node_module_available

        self.assertEqual("partial", result["status"])
        self.assertEqual("ready", result["routes"]["read"]["status"])
        self.assertEqual("ready", result["routes"]["edit_ooxml"]["status"])
        self.assertEqual("unavailable", result["routes"]["create"]["status"])
        self.assertEqual("unavailable", result["routes"]["visual_qa"]["status"])
        self.assertEqual("ready", read_only["status"])
        self.assertEqual(["read"], read_only["ready_routes"])
        self.assertEqual([], read_only["unavailable_routes"])
        self.assertEqual(
            {"python", "bundled_ooxml_reader"},
            set(read_only["capabilities"]),
        )
        self.assertEqual(str(Path.cwd().resolve()), read_only["execution_plan"]["dependency_root"])
        self.assertEqual(str(Path(sys.executable).resolve()), read_only["execution_plan"]["python"])
        self.assertEqual("3.9", read_only["capabilities"]["python"]["minimum_version"])

    def test_doctor_does_not_recommend_missing_unused_create_alternative(self):
        original_command_path = self.doctor.command_path
        original_libreoffice_path = self.doctor.libreoffice_path
        original_module_available = self.doctor.module_available
        original_node_module_available = self.doctor.node_module_available
        try:
            self.doctor.command_path = lambda *_names: None
            self.doctor.libreoffice_path = lambda: None
            self.doctor.module_available = lambda name: name == "docx"
            self.doctor.node_module_available = lambda _name, _root=None: False
            result = self.doctor.inspect("create")
        finally:
            self.doctor.command_path = original_command_path
            self.doctor.libreoffice_path = original_libreoffice_path
            self.doctor.module_available = original_module_available
            self.doctor.node_module_available = original_node_module_available

        self.assertEqual("ready", result["routes"]["create"]["status"])
        self.assertEqual({}, result["install_hints"])


class PptScriptTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.validator = load_script("ppt", "validate_pptx.py", "awaken_ppt_validate_test")
        cls.package = load_script("ppt", "ooxml_package.py", "awaken_ppt_package_test")
        cls.doctor = load_script("ppt", "doctor.py", "awaken_ppt_doctor_test")

    def test_valid_pptx_and_orphan_slide(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            valid = root / "valid.pptx"
            write_package(valid, minimal_pptx_files())
            self.assertEqual([], self.validator.inspect(valid))

            orphan_files = minimal_pptx_files()
            orphan_files["ppt/slides/slide2.xml"] = orphan_files["ppt/slides/slide1.xml"]
            orphan = root / "orphan.pptx"
            write_package(orphan, orphan_files)
            self.assertTrue(
                any("孤立 slide" in issue for issue in self.validator.inspect(orphan))
            )

    def test_stacked_chart_out_end_is_rejected(self):
        chart = """<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:plotArea><c:barChart><c:grouping val="stacked"/><c:dLbls><c:dLblPos val="outEnd"/></c:dLbls><c:axId val="1"/><c:axId val="2"/></c:barChart><c:catAx><c:axId val="1"/></c:catAx><c:valAx><c:axId val="2"/></c:valAx></c:plotArea></c:chart></c:chartSpace>"""
        from xml.etree import ElementTree

        issues = self.validator.chart_issues("ppt/charts/chart1.xml", ElementTree.fromstring(chart))
        self.assertTrue(any("outEnd" in issue for issue in issues))

    def test_doctor_scopes_structural_qa_without_visual_dependencies(self):
        original_command_path = self.doctor.command_path
        original_libreoffice_path = self.doctor.libreoffice_path
        original_module_available = self.doctor.module_available
        original_node_module_available = self.doctor.node_module_available
        try:
            self.doctor.command_path = lambda *_names: None
            self.doctor.libreoffice_path = lambda: None
            self.doctor.module_available = lambda _name: False
            self.doctor.node_module_available = lambda _name, _root=None: False
            result = self.doctor.inspect("structural_qa")
        finally:
            self.doctor.command_path = original_command_path
            self.doctor.libreoffice_path = original_libreoffice_path
            self.doctor.module_available = original_module_available
            self.doctor.node_module_available = original_node_module_available

        self.assertEqual("ready", result["status"])
        self.assertEqual(["structural_qa"], result["ready_routes"])
        self.assertEqual({"python", "bundled_validator"}, set(result["capabilities"]))


@unittest.skipUnless(importlib.util.find_spec("openpyxl"), "需要 openpyxl")
class ExcelScriptTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.inspector = load_script("excel", "inspect_workbook.py", "awaken_excel_inspect_test")
        cls.recalc = load_script("excel", "recalc.py", "awaken_excel_recalc_test")
        cls.doctor = load_script("excel", "doctor.py", "awaken_excel_doctor_test")

    def test_formula_and_dynamic_function_risk(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "model.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 10
            sheet["A2"] = "=SUM(A1:A1)"
            sheet["A3"] = '=XLOOKUP(A1,B1:B2,C1:C2)'
            workbook.save(path)
            result = self.inspector.inspect(path)
            self.assertEqual(2, result["formula_count"])
            self.assertTrue(any(item["function"] == "XLOOKUP" for item in result["risky_formulas"]))

    def test_inspector_profiles_headers_mixed_types_and_zero_values(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "mixed.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["金额", "状态"])
            sheet.append([0, "待处理"])
            sheet.append([12.5, "完成"])
            sheet.append(["已提交复核", "完成"])
            workbook.save(path)

            profile = self.inspector.inspect(path)["sheets"][0]["data_profile"]
            self.assertEqual(1, profile["header_candidates"][0]["row"])
            self.assertIn("A", profile["mixed_type_columns"])
            column_a = next(item for item in profile["columns"] if item["column"] == "A")
            self.assertEqual(1, column_a["zero_count"])
            self.assertEqual(2, column_a["type_counts"]["number"])
            self.assertEqual(2, column_a["type_counts"]["text"])

    def test_external_link_part_blocks_recalculation_before_libreoffice(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "external.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=1+1"
            workbook.save(path)
            with zipfile.ZipFile(path, "a") as package:
                package.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
            formulas, external = self.recalc.workbook_risks(path)
            self.assertEqual(1, formulas)
            self.assertEqual(["xl/externalLinks/externalLink1.xml"], external)

    def test_doctor_scopes_recalc_and_returns_absolute_plan(self):
        result = self.doctor.inspect("recalc")
        self.assertEqual({"python", "openpyxl", "libreoffice", "bundled_recalculator"}, set(result["capabilities"]))
        self.assertEqual(str(Path(sys.executable).resolve()), result["execution_plan"]["python"])
        self.assertEqual("3.9", result["capabilities"]["python"]["minimum_version"])


class OfficePythonCompatibilityTests(unittest.TestCase):
    def interpreters(self) -> list[str]:
        candidates = [sys.executable, shutil.which("python3.9")]
        xcode_python = Path(
            "/Applications/Xcode.app/Contents/Developer/Library/Frameworks/"
            "Python3.framework/Versions/3.9/bin/python3.9"
        )
        if xcode_python.is_file():
            candidates.append(str(xcode_python))
        return list(dict.fromkeys(str(Path(path).resolve()) for path in candidates if path))

    def test_office_temp_directories_work_on_available_python_39_and_current(self):
        for interpreter in self.interpreters():
            version = subprocess.run(
                [interpreter, "-c", "import sys; print('.'.join(map(str, sys.version_info[:2])))"],
                capture_output=True,
                text=True,
                check=True,
            ).stdout.strip()
            if tuple(int(part) for part in version.split(".")) < (3, 9):
                continue
            for skill in ("word", "excel", "ppt"):
                script_dir = SYSTEM_SKILLS / skill / "scripts"
                code = (
                    "import pathlib,sys;"
                    f"sys.path.insert(0,{str(script_dir)!r});"
                    "from office_runtime import temporary_directory;"
                    "\nwith temporary_directory('awaken-office-test-') as path:\n"
                    " assert isinstance(path, pathlib.Path)\n"
                    " assert path.is_dir()\n"
                    "saved = path\n"
                    "assert not saved.exists()\n"
                )
                completed = subprocess.run(
                    [interpreter, "-B", "-c", code],
                    capture_output=True,
                    text=True,
                    check=False,
                )
                self.assertEqual(
                    0,
                    completed.returncode,
                    f"{skill} failed on Python {version}: {completed.stderr}",
                )

    def test_no_office_script_uses_python_310_only_cleanup_argument(self):
        for skill in ("word", "excel", "ppt"):
            for script in (SYSTEM_SKILLS / skill / "scripts").glob("*.py"):
                self.assertNotIn("ignore_cleanup_errors", script.read_text(encoding="utf-8"), str(script))


if __name__ == "__main__":
    unittest.main()
